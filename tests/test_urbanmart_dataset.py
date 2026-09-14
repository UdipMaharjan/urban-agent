import json
from datetime import UTC, datetime
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
DATASET = ROOT / "data/raw/urbanmart_feedback.xlsx"
MANIFEST = ROOT / "data/reference/urbanmart_dataset_manifest.json"
HEADERS = ("feedback_id", "submitted_at", "rating", "feedback_text", "source")


def source_rows():
    manifest = json.loads(MANIFEST.read_text(encoding="utf-8"))
    assert manifest["expected_rows"] == 30
    assert len(manifest["records"]) == 30
    return manifest["records"]


def test_complete_urbanmart_workbook_preserves_source_data():
    expected = source_rows()
    workbook = load_workbook(BytesIO(DATASET.read_bytes()), read_only=True, data_only=False)
    assert workbook.sheetnames == ["Feedback"]
    sheet = workbook.active
    sheet.reset_dimensions()
    rows = list(sheet.iter_rows(values_only=True))
    workbook.close()

    assert rows[0] == HEADERS
    actual = []
    for feedback_id, submitted_at, rating, feedback_text, source in rows[1:]:
        actual.append(
            {
                "feedback_id": feedback_id,
                "submitted_at": submitted_at.strftime("%d-%b-%Y"),
                "rating": rating,
                "feedback_text": feedback_text,
                "source": source,
            }
        )
    assert actual == expected
    expected_ids = [f"F{number:03d}" for number in range(1, 31)]
    actual_ids = [record["feedback_id"] for record in actual]
    assert actual_ids == expected_ids
    assert len(actual_ids) == len(set(actual_ids)) == 30


def test_complete_urbanmart_dataset_import_and_duplicate_prevention(client, upload):
    expected = source_rows()
    response = upload(DATASET.read_bytes(), DATASET.name)
    assert response.status_code == 200
    assert response.json() == {
        "total_rows": 30,
        "imported_rows": 30,
        "skipped_rows": 0,
        "failed_rows": 0,
        "errors": [],
    }

    records = client.get("/api/feedback?limit=100").json()
    assert len(records) == 30
    for stored, original in zip(records, expected, strict=True):
        expected_date = datetime.strptime(original["submitted_at"], "%d-%b-%Y").replace(tzinfo=UTC)
        assert stored["feedback_id"] == original["feedback_id"]
        assert stored["submitted_at"] == expected_date.isoformat().replace("+00:00", "Z")
        assert stored["rating"] == original["rating"]
        assert stored["feedback_text"] == original["feedback_text"]
        assert stored["source"] == original["source"]

    duplicate = upload(DATASET.read_bytes(), DATASET.name).json()
    assert duplicate["total_rows"] == duplicate["skipped_rows"] == 30
    assert duplicate["imported_rows"] == duplicate["failed_rows"] == 0
    assert all(issue["code"] == "duplicate" for issue in duplicate["errors"])
    assert len(client.get("/api/feedback?limit=100").json()) == 30


def test_urbanmart_shaped_malformed_row_has_useful_error(upload, make_excel):
    content = make_excel(
        [["F999", datetime(2026, 9, 10), 2, None, "Online Survey"]],
        headers=list(HEADERS),
    )
    response = upload(content)
    assert response.status_code == 200
    summary = response.json()
    assert summary["total_rows"] == summary["failed_rows"] == 1
    assert summary["imported_rows"] == summary["skipped_rows"] == 0
    assert summary["errors"][0]["row"] == 2
    assert summary["errors"][0]["code"] == "validation_error"
    assert "feedback_text" in summary["errors"][0]["message"]
