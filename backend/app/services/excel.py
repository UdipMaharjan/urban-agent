from io import BytesIO
from zipfile import ZipFile

from openpyxl import load_workbook
from pydantic import ValidationError
from sqlalchemy.orm import Session

from ..config import Settings
from ..schemas import FeedbackCreate, ImportIssue, ImportSummary
from .feedback import store_feedback


class InvalidWorkbook(ValueError):
    pass


def read_rows(content: bytes, settings: Settings) -> list[tuple[int, dict, str | None]]:
    """Validate the whole file structure before any database mutation."""
    workbook = None
    try:
        with ZipFile(BytesIO(content)) as archive:
            if sum(info.file_size for info in archive.infolist()) > (
                settings.max_excel_uncompressed_bytes
            ):
                raise InvalidWorkbook("Workbook expands beyond the configured size limit.")
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=False)
        if len(workbook.sheetnames) != 1:
            raise InvalidWorkbook("Use exactly one worksheet so no sheets are silently ignored.")
        sheet = workbook.active
        # Do not trust producer-supplied dimensions: scan actual worksheet cells.
        sheet.reset_dimensions()
        iterator = sheet.iter_rows()
        header_cells = next(iterator, ())
        headers = [
            cell.value.strip() if isinstance(cell.value, str) else cell.value
            for cell in header_cells
        ]
        while headers and headers[-1] is None:
            headers.pop()
        if not headers or any(not isinstance(h, str) or not h for h in headers):
            raise InvalidWorkbook("Row 1 must contain non-empty column names.")
        if len(headers) != len(set(headers)):
            raise InvalidWorkbook("Duplicate column names are not allowed.")
        missing = {"feedback_id", "feedback_text"} - set(headers)
        if missing:
            raise InvalidWorkbook(f"Missing required columns: {', '.join(sorted(missing))}.")
        unknown = set(headers) - set(FeedbackCreate.model_fields)
        if unknown:
            raise InvalidWorkbook(f"Unknown columns: {', '.join(sorted(unknown))}.")
        rows = []
        for row_number, cells in enumerate(iterator, start=2):
            if row_number - 1 > settings.max_import_rows:
                raise InvalidWorkbook(f"Maximum {settings.max_import_rows} data rows per upload.")
            if any(cell.value is not None for cell in cells[len(headers) :]):
                raise InvalidWorkbook(f"Row {row_number} contains data without a column header.")
            issue = None
            if any(cell.data_type in {"f", "e"} for cell in cells):
                issue = "Formula and Excel error cells are not supported. Supply literal values."
            values = [cell.value for cell in cells[: len(headers)]]
            values += [None] * (len(headers) - len(values))
            payload = {
                key: value
                for key, value in zip(headers, values, strict=True)
                if value is not None and not (isinstance(value, str) and not value.strip())
            }
            rows.append((row_number, payload, issue))
        if not rows:
            raise InvalidWorkbook("Workbook contains headers but no data rows.")
        return rows
    except InvalidWorkbook:
        raise
    except Exception as exc:
        # Parser errors vary by malformed XML/ZIP. Never expose internal parser details.
        raise InvalidWorkbook(
            "Cannot read workbook. Upload a valid, unencrypted .xlsx file."
        ) from exc
    finally:
        if workbook is not None:
            workbook.close()


def import_excel(content: bytes, session: Session, settings: Settings) -> ImportSummary:
    rows = read_rows(content, settings)
    summary = ImportSummary(total_rows=len(rows))
    for row_number, values, issue in rows:
        source_id = values.get("feedback_id")
        source_id = str(source_id) if source_id is not None else None
        code = None
        message = ""
        if not values and issue is None:
            summary.skipped_rows += 1
            code, message = "blank_row", "Blank row skipped."
        else:
            try:
                if issue:
                    raise ValueError(issue)
                payload = FeedbackCreate.model_validate(values)
                _, outcome = store_feedback(session, payload)
                if outcome == "created":
                    summary.imported_rows += 1
                elif outcome == "duplicate":
                    summary.skipped_rows += 1
                    code, message = "duplicate", "Identical source and feedback_id already stored."
                else:
                    summary.failed_rows += 1
                    code, message = (
                        "conflict",
                        (
                            "This source and feedback_id already exist with different values. "
                            "Existing feedback was not changed."
                        ),
                    )
            except (ValidationError, ValueError) as exc:
                summary.failed_rows += 1
                code = "validation_error"
                if isinstance(exc, ValidationError):
                    message = "; ".join(
                        f"{'.'.join(map(str, error['loc']))}: {error['msg']}"
                        for error in exc.errors(include_input=False, include_url=False)
                    )
                else:
                    message = str(exc)
        if code:
            summary.errors.append(
                ImportIssue(row=row_number, feedback_id=source_id, code=code, message=message)
            )
    session.commit()
    return summary
